"""Tests for the Case API (Day 2), the legal-section loader (Day 3), and the
Seizure Receipt generation endpoint (Day 4).

Run: docker compose exec web python manage.py test core
"""
import tempfile
from pathlib import Path

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from .models import (
    Case,
    CaseSection,
    DiaryEvent,
    Document,
    LegalSection,
    Party,
    SeizedItem,
    Statement,
    User,
)

_TEST_MEDIA = tempfile.mkdtemp(prefix="crimegpt-media-")


class NestedCaseAPITests(APITestCase):
    def setUp(self):
        self.io = User.objects.create_user(
            username="io1", password="pw", role="IO", police_station="PS-Central"
        )
        self.client.force_authenticate(user=self.io)
        self.list_url = reverse("case-list")

    def _complete_payload(self):
        """A whole case in one request: 2 parties, 1 seized item, 1 statement."""
        return {
            "case_number": "C-001",
            "fir_number": "FIR-001",
            "fir_date": "2026-06-01",
            "police_station": "PS-Central",
            "status": "INVESTIGATION",
            "incident_place": "Market Road",
            "incident_narrative": "Mobile phone snatched.",
            "parties": [
                {"role": "VICTIM", "name": "Asha", "phone": "9990001111"},
                {"role": "ACCUSED", "name": "Ravi", "parent_name": "Mohan"},
            ],
            "seized_items": [
                {
                    "description": "One mobile phone",
                    "quantity": "1",
                    "seized_from_index": 1,  # -> the accused (Ravi)
                    "seizure_place": "Market Road",
                },
            ],
            "statements": [
                {
                    "party_index": 0,  # -> the victim (Asha)
                    "text": "I saw the accused snatch my phone.",
                    "recorded_date": "2026-06-02",
                },
            ],
        }

    # ------------------------------------------------------------------ #
    # CREATE
    # ------------------------------------------------------------------ #
    def test_create_complete_case_in_one_request(self):
        resp = self.client.post(self.list_url, self._complete_payload(), format="json")
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED, resp.data)

        # --- persisted correctly in the DB ---
        self.assertEqual(Case.objects.count(), 1)
        case = Case.objects.get()
        self.assertEqual(case.io, self.io)  # defaulted to the requesting user
        self.assertEqual(case.parties.count(), 2)
        self.assertEqual(case.seized_items.count(), 1)
        self.assertEqual(case.statements.count(), 1)

        victim = case.parties.get(role="VICTIM")
        accused = case.parties.get(role="ACCUSED")
        # index references resolved to the right parties
        self.assertEqual(case.seized_items.get().seized_from, accused)
        self.assertEqual(case.statements.get().party, victim)

        # --- reads back correctly through the API ---
        detail_url = reverse("case-detail", args=[case.id])
        data = self.client.get(detail_url).data
        self.assertEqual(data["case_number"], "C-001")
        self.assertEqual(data["io"], self.io.id)
        self.assertEqual(data["io_username"], "io1")
        self.assertEqual(len(data["parties"]), 2)
        self.assertEqual(len(data["seized_items"]), 1)
        self.assertEqual(len(data["statements"]), 1)
        self.assertEqual(data["seized_items"][0]["seized_from"], accused.id)
        self.assertEqual(data["statements"][0]["party"], victim.id)
        names = {p["name"] for p in data["parties"]}
        self.assertEqual(names, {"Asha", "Ravi"})

    def test_bad_party_index_is_rejected_and_rolls_back(self):
        payload = self._complete_payload()
        payload["statements"][0]["party_index"] = 5  # out of range
        resp = self.client.post(self.list_url, payload, format="json")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        # transaction.atomic must roll the half-written case back
        self.assertEqual(Case.objects.count(), 0)

    # ------------------------------------------------------------------ #
    # UPDATE
    # ------------------------------------------------------------------ #
    def test_update_replaces_nested_items(self):
        case_id = self.client.post(
            self.list_url, self._complete_payload(), format="json"
        ).data["id"]

        update = {
            "case_number": "C-001",
            "fir_number": "FIR-001",
            "fir_date": "2026-06-01",
            "police_station": "PS-Central",
            "status": "ARREST",  # changed
            "incident_place": "Market Road",
            "incident_narrative": "Mobile phone snatched. Accused arrested.",
            "parties": [
                {"role": "VICTIM", "name": "Asha", "phone": "9990001111"},
                {"role": "ACCUSED", "name": "Ravi Kumar", "parent_name": "Mohan"},  # renamed
                {"role": "WITNESS", "name": "Sunil"},  # added
            ],
            "seized_items": [
                {"description": "One mobile phone (black)", "quantity": "1", "seized_from_index": 1},
                {"description": "Cash", "quantity": "Rs 500", "seized_from_index": 1},  # added
            ],
            "statements": [
                {"party_index": 2, "text": "I witnessed the arrest.", "recorded_date": "2026-06-03"},
            ],
        }
        detail_url = reverse("case-detail", args=[case_id])
        resp = self.client.put(detail_url, update, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.data)

        case = Case.objects.get(id=case_id)
        self.assertEqual(case.status, "ARREST")
        self.assertEqual(case.parties.count(), 3)       # 2 -> 3
        self.assertEqual(case.seized_items.count(), 2)  # 1 -> 2
        self.assertEqual(case.statements.count(), 1)

        accused = case.parties.get(role="ACCUSED")
        witness = case.parties.get(role="WITNESS")
        self.assertEqual(accused.name, "Ravi Kumar")              # nested update applied
        self.assertEqual(case.statements.get().party, witness)   # index 2 -> new witness
        self.assertEqual(
            set(case.seized_items.values_list("description", flat=True)),
            {"One mobile phone (black)", "Cash"},
        )
        for item in case.seized_items.all():
            self.assertEqual(item.seized_from, accused)

    def test_partial_update_leaves_nested_untouched(self):
        case_id = self.client.post(
            self.list_url, self._complete_payload(), format="json"
        ).data["id"]
        detail_url = reverse("case-detail", args=[case_id])

        # PATCH only a scalar field; omit nested keys entirely.
        resp = self.client.patch(detail_url, {"status": "CHARGESHEET"}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.data)

        case = Case.objects.get(id=case_id)
        self.assertEqual(case.status, "CHARGESHEET")
        # nested collections preserved
        self.assertEqual(case.parties.count(), 2)
        self.assertEqual(case.seized_items.count(), 1)
        self.assertEqual(case.statements.count(), 1)

    # ------------------------------------------------------------------ #
    # DELETE
    # ------------------------------------------------------------------ #
    def test_delete_case_cascades(self):
        case_id = self.client.post(
            self.list_url, self._complete_payload(), format="json"
        ).data["id"]
        resp = self.client.delete(reverse("case-detail", args=[case_id]))
        self.assertEqual(resp.status_code, status.HTTP_204_NO_CONTENT)
        self.assertEqual(Case.objects.count(), 0)


class LoadLegalSectionsCommandTests(TestCase):
    """Tests for the idempotent CSV -> LegalSection loader (Day 3)."""

    HEADER = "code,section_number,title,text,ipc_equivalent\n"

    def _write_csv(self, body):
        tmp = tempfile.NamedTemporaryFile(
            mode="w", suffix=".csv", delete=False, encoding="utf-8", newline=""
        )
        tmp.write(self.HEADER + body)
        tmp.close()
        self.addCleanup(lambda: Path(tmp.name).unlink(missing_ok=True))
        return tmp.name

    def test_loads_rows_from_csv(self):
        path = self._write_csv(
            'BNS,103,Punishment for murder.,Whoever commits murder...,302\n'
            'BNS,303,Theft.,Whoever intends to take...,"378, 379"\n'
        )
        call_command("load_legal_sections", path)

        self.assertEqual(LegalSection.objects.count(), 2)
        murder = LegalSection.objects.get(code="BNS", section_number="103")
        self.assertEqual(murder.title, "Punishment for murder.")
        self.assertEqual(murder.ipc_equivalent, "302")
        theft = LegalSection.objects.get(code="BNS", section_number="303")
        self.assertEqual(theft.ipc_equivalent, "378, 379")

    def test_rerun_is_idempotent(self):
        path = self._write_csv(
            "BNS,103,Punishment for murder.,text,302\n"
            "BNS,303,Theft.,text,378\n"
        )
        call_command("load_legal_sections", path)
        call_command("load_legal_sections", path)  # second run
        # No duplicates created on re-run.
        self.assertEqual(LegalSection.objects.count(), 2)

    def test_rerun_updates_changed_rows(self):
        first = self._write_csv("BNS,103,Old title,old text,302\n")
        call_command("load_legal_sections", first)

        second = self._write_csv("BNS,103,Punishment for murder.,new text,302\n")
        call_command("load_legal_sections", second)

        self.assertEqual(LegalSection.objects.count(), 1)  # updated in place
        murder = LegalSection.objects.get(code="BNS", section_number="103")
        self.assertEqual(murder.title, "Punishment for murder.")
        self.assertEqual(murder.text, "new text")

    def test_missing_columns_raises(self):
        bad = tempfile.NamedTemporaryFile(
            mode="w", suffix=".csv", delete=False, encoding="utf-8", newline=""
        )
        bad.write("code,section_number,title\nBNS,103,Murder\n")  # missing text, ipc
        bad.close()
        self.addCleanup(lambda: Path(bad.name).unlink(missing_ok=True))
        with self.assertRaises(CommandError):
            call_command("load_legal_sections", bad.name)

    def test_dry_run_writes_nothing(self):
        path = self._write_csv("BNS,103,Murder,text,302\n")
        call_command("load_legal_sections", path, "--dry-run")
        self.assertEqual(LegalSection.objects.count(), 0)


class SeizureReceiptTemplateTests(TestCase):
    """Day 3 part 2: the Seizure Receipt template renders with sample data."""

    def _all_text(self, docx_path):
        from docx import Document

        doc = Document(docx_path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)

    def test_template_renders_with_sample_data(self):
        import tempfile

        from core.documents import (
            build_seizure_receipt_template,
            render_seizure_receipt,
            sample_seizure_context,
        )

        # Build the template fresh (also exercises the builder), then render.
        tpl_path = Path(tempfile.mkdtemp()) / "seizure_receipt.docx"
        build_seizure_receipt_template(tpl_path)
        self.assertTrue(tpl_path.exists())

        out_path = Path(tempfile.mkdtemp()) / "rendered.docx"
        render_seizure_receipt(sample_seizure_context(), template_path=tpl_path).save(
            str(out_path)
        )

        text = self._all_text(out_path)
        # Scalar fields filled in
        self.assertIn("C-2026-001", text)
        self.assertIn("FIR-045/2026", text)
        self.assertIn("Inspector R. Sharma", text)
        self.assertIn("Ravi Kumar", text)
        # Both seized items rendered (table loop worked)
        self.assertIn("One black mobile phone (Samsung)", text)
        self.assertIn("Rs. 5,000", text)
        # Witness loop worked
        self.assertIn("Sunil Verma", text)
        self.assertIn("Asha Devi", text)
        # No unrendered Jinja tags left behind
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)


@override_settings(MEDIA_ROOT=_TEST_MEDIA)
class SeizureGenerationAPITests(APITestCase):
    """Day 4: POST /api/cases/{id}/documents/seizure/."""

    def setUp(self):
        self.io = User.objects.create_user(username="io1", password="pw", role="IO", police_station="PS-Central")
        self.sho = User.objects.create_user(username="sho1", password="pw", role="SHO", police_station="PS-Central")
        self.legal = User.objects.create_user(username="legal1", password="pw", role="LEGAL")

        self.case = Case.objects.create(
            case_number="C-900", fir_number="FIR-900", fir_date="2026-06-01",
            police_station="PS-Central", status="INVESTIGATION",
            incident_place="Market Road", io=self.io,
        )
        self.accused = Party.objects.create(case=self.case, role="ACCUSED", name="Ravi", parent_name="Mohan")
        Party.objects.create(case=self.case, role="WITNESS", name="Sunil", address="5 Nehru Rd")
        SeizedItem.objects.create(
            case=self.case, description="One mobile phone", quantity="1",
            seized_from=self.accused, seizure_place="Market Road",
        )
        self.url = reverse("case-generate-seizure", args=[self.case.id])

    def test_io_can_generate_receipt(self):
        self.client.force_authenticate(self.io)
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED, resp.data)
        self.assertEqual(resp.data["version"], 1)
        self.assertTrue(resp.data["download_url"].endswith(".docx"))

        self.assertEqual(Document.objects.filter(case=self.case, doc_type="SEIZURE").count(), 1)
        doc = Document.objects.get(case=self.case, doc_type="SEIZURE")
        self.assertTrue(doc.file.name)

        # The rendered file actually contains the case data.
        from docx import Document as Docx

        doc.file.open("rb")
        try:
            rendered = Docx(doc.file)
        finally:
            doc.file.close()
        text = "\n".join(p.text for p in rendered.paragraphs)
        for table in rendered.tables:
            for row in table.rows:
                text += "\n" + " | ".join(c.text for c in row.cells)
        self.assertIn("C-900", text)
        self.assertIn("Ravi", text)
        self.assertIn("One mobile phone", text)

    def test_sho_can_generate_receipt(self):
        self.client.force_authenticate(self.sho)
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED, resp.data)

    def test_regeneration_increments_version(self):
        self.client.force_authenticate(self.io)
        first = self.client.post(self.url)
        second = self.client.post(self.url)
        self.assertEqual(first.data["version"], 1)
        self.assertEqual(second.data["version"], 2)
        # Old version is kept, not overwritten.
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="SEIZURE").count(), 2)
        self.assertEqual(
            sorted(Document.objects.filter(case=self.case).values_list("version", flat=True)),
            [1, 2],
        )

    def test_diary_event_is_logged(self):
        self.client.force_authenticate(self.io)
        self.assertEqual(self.case.diary.count(), 0)
        self.client.post(self.url)

        events = self.case.diary.filter(event_type="DOCUMENT_GENERATED")
        self.assertEqual(events.count(), 1)
        event = events.get()
        self.assertEqual(event.actor, self.io)
        self.assertIn("Seizure Receipt", event.description)

    def test_legal_advisor_is_forbidden(self):
        self.client.force_authenticate(self.legal)
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        # Nothing was generated or logged.
        self.assertEqual(Document.objects.filter(case=self.case).count(), 0)
        self.assertEqual(self.case.diary.count(), 0)

    def test_unauthenticated_is_rejected(self):
        resp = self.client.post(self.url)
        self.assertIn(resp.status_code, (status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN))
        self.assertEqual(Document.objects.filter(case=self.case).count(), 0)


class SeizureContextMappingTests(TestCase):
    """Day-5 fix: the seized-from person must appear in the narrative AND
    signature lines, not only in the items table."""

    def _paragraph_and_table_text(self, case):
        """Render and return (paragraph_text, table_text) separately.

        The narrative and signature are *paragraphs*; the items table is *table
        cells*. python-docx keeps these separate, so asserting against the
        paragraph text alone proves the name reached the narrative/signature.
        """
        from io import BytesIO

        from docx import Document as Docx

        from core.documents import build_seizure_context_from_case, render_seizure_receipt

        tpl = render_seizure_receipt(build_seizure_context_from_case(case))
        buf = BytesIO()
        tpl.save(buf)
        buf.seek(0)
        doc = Docx(buf)
        paragraph_text = "\n".join(p.text for p in doc.paragraphs)
        table_text = "\n".join(
            c.text for t in doc.tables for r in t.rows for c in r.cells
        )
        return paragraph_text, table_text

    def _make_case(self, number):
        io = User.objects.create_user(username=f"io_{number}", password="pw", role="IO")
        return Case.objects.create(
            case_number=number, fir_number="FIR-1", fir_date="2026-06-01",
            police_station="PS-Central", incident_place="Market Road", io=io,
        )

    def test_accused_name_in_narrative_and_signature(self):
        case = self._make_case("MAP-1")
        accused = Party.objects.create(
            case=case, role="ACCUSED", name="Ravi Kumar",
            parent_name="Mohan Kumar", address="12 MG Road",
        )
        SeizedItem.objects.create(case=case, description="Phone", seized_from=accused)

        paragraphs, _ = self._paragraph_and_table_text(case)
        # Name, parent and address all reach the narrative/signature paragraphs.
        self.assertIn("Ravi Kumar", paragraphs)
        self.assertIn("Mohan Kumar", paragraphs)
        self.assertIn("12 MG Road", paragraphs)
        # And specifically the narrative + signature lines, not a blank fragment.
        self.assertIn("seized the following article(s)", paragraphs)
        self.assertNotIn("from , child/spouse of , resident of :", paragraphs)
        self.assertNotIn("from whom seized: __________________  ()", paragraphs)

    def test_falls_back_to_seized_from_when_no_accused(self):
        """No ACCUSED party: the person items were seized from should be named."""
        case = self._make_case("MAP-2")
        victim = Party.objects.create(
            case=case, role="VICTIM", name="Asha Devi",
            parent_name="Ram Devi", address="9 Lake View",
        )
        SeizedItem.objects.create(case=case, description="Bag", seized_from=victim)
        self.assertFalse(case.parties.filter(role="ACCUSED").exists())

        paragraphs, table = self._paragraph_and_table_text(case)
        # The fallback person appears in the narrative/signature paragraphs...
        self.assertIn("Asha Devi", paragraphs)
        self.assertIn("Ram Devi", paragraphs)
        self.assertIn("9 Lake View", paragraphs)
        # ...not merely in the items table.
        self.assertIn("Asha Devi", table)
        self.assertNotIn("from , child/spouse of , resident of :", paragraphs)


class LegalRetrievalTests(APITestCase):
    """Week 2 Track A: local embedding + pgvector section suggestions."""

    # A small representative corpus with realistic BNS-style text.
    CORPUS = [
        ("303", "Theft.",
         "Whoever, intending to take dishonestly any movable property out of the "
         "possession of any person without that person's consent, moves that "
         "property in order to such taking, is said to commit theft."),
        ("305", "Theft in a dwelling house.",
         "Whoever commits theft in any building, tent or vessel used as a human "
         "dwelling or for the custody of property shall be punished with "
         "imprisonment and fine."),
        ("330", "House-trespass and house-breaking.",
         "Whoever commits criminal trespass by entering into or remaining in any "
         "building used as a human dwelling commits house-trespass; house-breaking "
         "is committed by effecting entrance into a house by breaking a lock or door."),
        ("103", "Punishment for murder.",
         "Whoever commits murder shall be punished with death or imprisonment for "
         "life, and shall also be liable to fine."),
        ("318", "Cheating.",
         "Whoever by deceiving any person fraudulently or dishonestly induces the "
         "person so deceived to deliver any property is said to cheat."),
        ("309", "Robbery.",
         "In all robbery there is either theft or extortion; theft is robbery if "
         "the offender voluntarily causes or attempts to cause death, hurt or "
         "wrongful restraint while committing the theft."),
        ("351", "Criminal intimidation.",
         "Whoever threatens another with any injury to his person, reputation or "
         "property with intent to cause alarm commits criminal intimidation."),
        ("115", "Voluntarily causing hurt.",
         "Whoever causes bodily pain, disease or infirmity to any person voluntarily "
         "is said to voluntarily cause hurt and shall be punished."),
    ]

    THEFT_NARRATIVE = (
        "The accused broke open the lock of the complainant's house at night and "
        "stole gold ornaments and cash from inside the dwelling."
    )

    def setUp(self):
        for number, title, text in self.CORPUS:
            LegalSection.objects.create(
                code="BNS", section_number=number, title=title, text=text
            )
        self.io = User.objects.create_user(username="io_ret", password="pw", role="IO")
        self.case = Case.objects.create(
            case_number="RET-1", fir_number="FIR-RET", fir_date="2026-06-01",
            police_station="PS-Central", incident_narrative=self.THEFT_NARRATIVE, io=self.io,
        )
        self.client.force_authenticate(self.io)
        self.url = reverse("case-suggest-sections", args=[self.case.id])

    # -- embedding command ------------------------------------------------
    def test_embed_command_populates_vectors(self):
        self.assertEqual(LegalSection.objects.filter(embedding__isnull=False).count(), 0)
        call_command("embed_legal_sections")
        sections = LegalSection.objects.all()
        self.assertTrue(all(s.embedding is not None for s in sections))
        self.assertEqual(len(sections[0].embedding), 384)

    def test_embed_command_is_idempotent(self):
        from io import StringIO

        call_command("embed_legal_sections")
        out = StringIO()
        call_command("embed_legal_sections", stdout=out)  # second run
        self.assertIn("Nothing to embed", out.getvalue())

    # -- suggestion endpoint ---------------------------------------------
    def test_suggest_returns_sensible_ranked_sections(self):
        call_command("embed_legal_sections")
        resp = self.client.post(self.url, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.data)

        results = resp.data["results"]
        self.assertGreaterEqual(len(results), 1)
        self.assertLessEqual(len(results), 5)  # default top-N

        # Scores are descending.
        scores = [r["similarity"] for r in results]
        self.assertEqual(scores, sorted(scores, reverse=True))

        # The top hit is theft/house-trespass, not murder/cheating/hurt.
        top_title = results[0]["title"].lower()
        self.assertTrue(
            "theft" in top_title or "trespass" in top_title or "house" in top_title,
            f"unexpected top section: {results[0]}",
        )
        top3 = {r["section_number"] for r in results[:3]}
        self.assertTrue(top3 & {"303", "305", "330"}, f"top-3 were {top3}")

    def test_suggestions_are_created_unconfirmed(self):
        call_command("embed_legal_sections")
        resp = self.client.post(self.url, format="json")
        results = resp.data["results"]

        links = CaseSection.objects.filter(case=self.case)
        self.assertEqual(links.count(), len(results))
        for link in links:
            self.assertTrue(link.suggested_by_ai)
            self.assertFalse(link.confirmed)  # never auto-confirmed
        # response echoes the unconfirmed state
        self.assertTrue(all(r["suggested_by_ai"] and not r["confirmed"] for r in results))

    def test_rerun_does_not_duplicate_or_unconfirm(self):
        call_command("embed_legal_sections")
        first = self.client.post(self.url, format="json").data["results"]
        count_after_first = CaseSection.objects.filter(case=self.case).count()

        # Officer confirms one suggestion (the separate explicit action).
        confirmed_id = first[0]["case_section_id"]
        cs = CaseSection.objects.get(id=confirmed_id)
        cs.confirmed = True
        cs.save()

        # Re-run suggest.
        self.client.post(self.url, format="json")

        # No duplicate rows, and the confirmed one stays confirmed.
        self.assertEqual(
            CaseSection.objects.filter(case=self.case).count(), count_after_first
        )
        cs.refresh_from_db()
        self.assertTrue(cs.confirmed)

    def test_suggest_requires_a_narrative(self):
        self.case.incident_narrative = ""
        self.case.save()
        resp = self.client.post(self.url, format="json")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)


@override_settings(MEDIA_ROOT=_TEST_MEDIA)
class RemandRequestTests(APITestCase):
    """Week 2 Track B: Police-Custody Remand Request."""

    def setUp(self):
        from django.utils import timezone

        self.io = User.objects.create_user(username="io_rem", password="pw", role="IO", police_station="PS-Central")
        self.legal = User.objects.create_user(username="legal_rem", password="pw", role="LEGAL")
        self.case = Case.objects.create(
            case_number="REM-1", fir_number="FIR-REM", fir_date="2026-06-01",
            police_station="PS-Central", status="ARREST",
            incident_narrative="Accused broke into the dwelling and stole jewellery.",
            io=self.io,
        )
        Party.objects.create(
            case=self.case, role="ACCUSED", name="Ravi Kumar",
            parent_name="Mohan Kumar", address="12 MG Road",
        )
        # One CONFIRMED section and one AI-suggested-but-UNCONFIRMED section.
        self.confirmed_sec = LegalSection.objects.create(
            code="BNS", section_number="305", title="Theft in a dwelling house.", text="...")
        self.unconfirmed_sec = LegalSection.objects.create(
            code="BNS", section_number="318", title="Cheating.", text="...")
        CaseSection.objects.create(case=self.case, section=self.confirmed_sec, suggested_by_ai=True, confirmed=True)
        CaseSection.objects.create(case=self.case, section=self.unconfirmed_sec, suggested_by_ai=True, confirmed=False)
        # Arrest diary event (source of arrest details).
        DiaryEvent.objects.create(
            case=self.case, timestamp=timezone.make_aware(__import__("datetime").datetime(2026, 6, 2, 9, 30)),
            event_type="ARREST", description="Accused arrested at the railway station.", actor=self.io,
        )
        self.url = reverse("case-generate-remand", args=[self.case.id])

    def _paragraph_text(self, case):
        from io import BytesIO

        from docx import Document as Docx

        from core.documents import build_remand_context_from_case, render_remand_request

        case.refresh_from_db()  # ensure fir_date is a real date (as get_object() yields)
        tpl = render_remand_request(build_remand_context_from_case(case))
        buf = BytesIO()
        tpl.save(buf)
        buf.seek(0)
        return "\n".join(p.text for p in Docx(buf).paragraphs)

    def test_render_lands_real_data_and_only_confirmed_sections(self):
        text = self._paragraph_text(self.case)
        # Accused details in the narrative/signature paragraphs.
        self.assertIn("Ravi Kumar", text)
        self.assertIn("Mohan Kumar", text)
        self.assertIn("12 MG Road", text)
        # Confirmed section is present...
        self.assertIn("305", text)
        self.assertIn("Theft in a dwelling house", text)
        # ...and the UNCONFIRMED section is NOT (load-bearing boundary).
        self.assertNotIn("318", text)
        self.assertNotIn("Cheating", text)
        # Arrest details from the diary event.
        self.assertIn("railway station", text)
        self.assertIn("02-06-2026", text)
        # No unrendered Jinja.
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    def test_io_generates_versioned_with_diary(self):
        self.client.force_authenticate(self.io)
        r1 = self.client.post(self.url)
        self.assertEqual(r1.status_code, status.HTTP_201_CREATED, r1.data)
        self.assertEqual(r1.data["doc_type"], "REMAND")
        self.assertEqual(r1.data["version"], 1)
        self.assertTrue(r1.data["download_url"].endswith(".docx"))

        r2 = self.client.post(self.url)  # regenerate -> v2, not overwritten
        self.assertEqual(r2.data["version"], 2)
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="REMAND").count(), 2)

        events = self.case.diary.filter(event_type="DOCUMENT_GENERATED")
        self.assertEqual(events.count(), 2)
        self.assertIn("Remand Request", events.first().description)
        self.assertEqual(events.first().actor, self.io)

    def test_legal_advisor_forbidden(self):
        self.client.force_authenticate(self.legal)
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="REMAND").count(), 0)


@override_settings(MEDIA_ROOT=_TEST_MEDIA)
class CourtCustodyLetterTests(APITestCase):
    """Week 2 Track B: Court (Judicial) Custody Letter."""

    def setUp(self):
        self.io = User.objects.create_user(username="io_cc", password="pw", role="IO", police_station="PS-Central")
        self.legal = User.objects.create_user(username="legal_cc", password="pw", role="LEGAL")
        self.case = Case.objects.create(
            case_number="CC-1", fir_number="FIR-CC", fir_date="2026-06-01",
            police_station="PS-Central", status="ARREST",
            incident_narrative="Accused broke into the dwelling and stole jewellery.",
            io=self.io,
        )
        Party.objects.create(
            case=self.case, role="ACCUSED", name="Ravi Kumar",
            parent_name="Mohan Kumar", address="12 MG Road",
        )
        self.confirmed_sec = LegalSection.objects.create(
            code="BNS", section_number="305", title="Theft in a dwelling house.", text="...")
        self.unconfirmed_sec = LegalSection.objects.create(
            code="BNS", section_number="318", title="Cheating.", text="...")
        CaseSection.objects.create(case=self.case, section=self.confirmed_sec, suggested_by_ai=True, confirmed=True)
        CaseSection.objects.create(case=self.case, section=self.unconfirmed_sec, suggested_by_ai=True, confirmed=False)
        self.url = reverse("case-generate-custody", args=[self.case.id])

    def _paragraph_text(self, case):
        from io import BytesIO

        from docx import Document as Docx

        from core.documents import (
            build_court_custody_context_from_case,
            render_court_custody_letter,
        )

        case.refresh_from_db()  # ensure fir_date is a real date (as get_object() yields)
        tpl = render_court_custody_letter(build_court_custody_context_from_case(case))
        buf = BytesIO()
        tpl.save(buf)
        buf.seek(0)
        return "\n".join(p.text for p in Docx(buf).paragraphs)

    def test_render_lands_real_data_and_only_confirmed_sections(self):
        text = self._paragraph_text(self.case)
        self.assertIn("Ravi Kumar", text)
        self.assertIn("Mohan Kumar", text)
        self.assertIn("12 MG Road", text)
        # Confirmed section present...
        self.assertIn("305", text)
        self.assertIn("Theft in a dwelling house", text)
        # ...unconfirmed section absent (load-bearing boundary).
        self.assertNotIn("318", text)
        self.assertNotIn("Cheating", text)
        # No unrendered Jinja.
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    def test_io_generates_versioned_with_diary(self):
        self.client.force_authenticate(self.io)
        r1 = self.client.post(self.url)
        self.assertEqual(r1.status_code, status.HTTP_201_CREATED, r1.data)
        self.assertEqual(r1.data["doc_type"], "CUSTODY")
        self.assertEqual(r1.data["version"], 1)
        self.assertTrue(r1.data["download_url"].endswith(".docx"))

        r2 = self.client.post(self.url)  # regenerate -> v2
        self.assertEqual(r2.data["version"], 2)
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="CUSTODY").count(), 2)

        events = self.case.diary.filter(event_type="DOCUMENT_GENERATED")
        self.assertEqual(events.count(), 2)
        self.assertIn("Court Custody Letter", events.first().description)
        self.assertEqual(events.first().actor, self.io)

    def test_legal_advisor_forbidden(self):
        self.client.force_authenticate(self.legal)
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="CUSTODY").count(), 0)


@override_settings(MEDIA_ROOT=_TEST_MEDIA)
class MedicalTreatmentLetterTests(APITestCase):
    """Week 2 Track B: Medical Treatment Letter (subject VICTIM -> ACCUSED -> blank)."""

    def setUp(self):
        self.io = User.objects.create_user(username="io_med", password="pw", role="IO", police_station="PS-Central")
        self.legal = User.objects.create_user(username="legal_med", password="pw", role="LEGAL")

    def _make_case(self, number):
        return Case.objects.create(
            case_number=number, fir_number="FIR-MED", fir_date="2026-06-01",
            police_station="PS-Central",
            incident_narrative="Complainant was assaulted and injured.",
            io=self.io,
        )

    def _paragraph_text(self, case):
        from io import BytesIO

        from docx import Document as Docx

        from core.documents import build_medical_context_from_case, render_medical_treatment_letter

        tpl = render_medical_treatment_letter(build_medical_context_from_case(case))
        buf = BytesIO()
        tpl.save(buf)
        buf.seek(0)
        return "\n".join(p.text for p in Docx(buf).paragraphs)

    def test_subject_is_victim_when_present(self):
        case = self._make_case("MED-1")
        Party.objects.create(case=case, role="VICTIM", name="Asha Devi",
                             parent_name="Ram Devi", age=30, gender="Female", address="9 Lake View")
        Party.objects.create(case=case, role="ACCUSED", name="Ravi Kumar",
                             parent_name="Mohan", address="12 MG Road")

        text = self._paragraph_text(case)
        self.assertIn("Asha Devi", text)        # victim is the subject
        self.assertIn("9 Lake View", text)
        self.assertIn("Victim", text)           # role label rendered
        self.assertNotIn("Ravi Kumar", text)    # accused is NOT the subject
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    def test_falls_back_to_accused_when_no_victim(self):
        case = self._make_case("MED-2")
        Party.objects.create(case=case, role="ACCUSED", name="Ravi Kumar",
                             parent_name="Mohan Kumar", age=25, gender="Male", address="12 MG Road")
        self.assertFalse(case.parties.filter(role="VICTIM").exists())

        text = self._paragraph_text(case)
        self.assertIn("Ravi Kumar", text)       # accused used as fallback subject
        self.assertIn("12 MG Road", text)
        self.assertIn("Accused", text)
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    def test_blank_subject_when_no_victim_or_accused(self):
        case = self._make_case("MED-3")
        Party.objects.create(case=case, role="WITNESS", name="Sunil", address="5 Nehru Rd")

        text = self._paragraph_text(case)
        # Renders cleanly with a blank subject (no crash, no leftover tags),
        # and the witness is not pulled in as the subject.
        self.assertNotIn("Sunil", text)
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    def test_io_generates_versioned_with_diary(self):
        case = self._make_case("MED-4")
        Party.objects.create(case=case, role="VICTIM", name="Asha Devi", address="9 Lake View")
        url = reverse("case-generate-medical", args=[case.id])
        self.client.force_authenticate(self.io)

        r1 = self.client.post(url)
        self.assertEqual(r1.status_code, status.HTTP_201_CREATED, r1.data)
        self.assertEqual(r1.data["doc_type"], "MEDICAL")
        self.assertEqual(r1.data["version"], 1)

        r2 = self.client.post(url)
        self.assertEqual(r2.data["version"], 2)
        self.assertEqual(Document.objects.filter(case=case, doc_type="MEDICAL").count(), 2)
        events = case.diary.filter(event_type="DOCUMENT_GENERATED")
        self.assertEqual(events.count(), 2)
        self.assertIn("Medical Treatment Letter", events.first().description)

    def test_legal_advisor_forbidden(self):
        case = self._make_case("MED-5")
        Party.objects.create(case=case, role="VICTIM", name="Asha Devi", address="9 Lake View")
        url = reverse("case-generate-medical", args=[case.id])
        self.client.force_authenticate(self.legal)
        resp = self.client.post(url)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(Document.objects.filter(case=case, doc_type="MEDICAL").count(), 0)


@override_settings(MEDIA_ROOT=_TEST_MEDIA)
class ChargesheetTests(APITestCase):
    """Week 2 Track B: Purvani Chargesheet — the superset document."""

    def setUp(self):
        import datetime

        from django.utils import timezone

        self.io = User.objects.create_user(username="io_cs", password="pw", role="IO", police_station="PS-Central")
        self.legal = User.objects.create_user(username="legal_cs", password="pw", role="LEGAL")
        self.case = Case.objects.create(
            case_number="CS-1", fir_number="FIR-CS", fir_date="2026-06-01",
            police_station="PS-Central", status="CHARGESHEET",
            incident_place="Market Road",
            incident_narrative="Accused broke into the dwelling and stole jewellery.",
            io=self.io,
        )
        self.victim = Party.objects.create(case=self.case, role="VICTIM", name="Asha Devi", parent_name="Ram Devi", address="9 Lake View")
        self.accused = Party.objects.create(case=self.case, role="ACCUSED", name="Ravi Kumar", parent_name="Mohan Kumar", address="12 MG Road")
        self.witness = Party.objects.create(case=self.case, role="WITNESS", name="Sunil Verma", address="5 Nehru Road")

        self.confirmed_sec = LegalSection.objects.create(code="BNS", section_number="305", title="Theft in a dwelling house.", text="...")
        self.unconfirmed_sec = LegalSection.objects.create(code="BNS", section_number="318", title="Cheating.", text="...")
        CaseSection.objects.create(case=self.case, section=self.confirmed_sec, suggested_by_ai=True, confirmed=True)
        CaseSection.objects.create(case=self.case, section=self.unconfirmed_sec, suggested_by_ai=True, confirmed=False)

        SeizedItem.objects.create(case=self.case, description="One mobile phone", quantity="1", seized_from=self.accused, seizure_place="Market Road")
        SeizedItem.objects.create(case=self.case, description="Stolen gold chain", quantity="1", seized_from=self.accused)

        Statement.objects.create(case=self.case, party=self.victim, text="The accused broke into my house.", recorded_date="2026-06-02")
        Statement.objects.create(case=self.case, party=self.witness, text="I saw him climb the wall.", recorded_date="2026-06-02")

        DiaryEvent.objects.create(case=self.case, timestamp=timezone.make_aware(datetime.datetime(2026, 6, 1, 10, 0)), event_type="FIR_REGISTERED", description="FIR registered.", actor=self.io)
        DiaryEvent.objects.create(case=self.case, timestamp=timezone.make_aware(datetime.datetime(2026, 6, 2, 9, 30)), event_type="ARREST", description="Accused arrested at the railway station.", actor=self.io)

        self.url = reverse("case-generate-chargesheet", args=[self.case.id])

    def _all_text(self, case):
        from io import BytesIO

        from docx import Document as Docx

        from core.documents import build_chargesheet_context_from_case, render_chargesheet

        case.refresh_from_db()
        tpl = render_chargesheet(build_chargesheet_context_from_case(case))
        buf = BytesIO()
        tpl.save(buf)
        buf.seek(0)
        doc = Docx(buf)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)

    def test_render_includes_everything_and_only_confirmed_sections(self):
        text = self._all_text(self.case)

        # All party roles + names land.
        for name in ("Asha Devi", "Ravi Kumar", "Sunil Verma"):
            self.assertIn(name, text)
        for role in ("Victim", "Accused", "Witness"):
            self.assertIn(role, text)

        # Confirmed section present; unconfirmed absent (load-bearing boundary).
        self.assertIn("305", text)
        self.assertIn("Theft in a dwelling house", text)
        self.assertNotIn("318", text)
        self.assertNotIn("Cheating", text)

        # Seized items land.
        self.assertIn("One mobile phone", text)
        self.assertIn("Stolen gold chain", text)

        # Statements land (deponent + text).
        self.assertIn("The accused broke into my house.", text)
        self.assertIn("I saw him climb the wall.", text)

        # Diary summary lands.
        self.assertIn("ARREST", text)
        self.assertIn("railway station", text)

        # Brief facts.
        self.assertIn("broke into the dwelling", text)

        # No unrendered Jinja anywhere (tables + paragraphs).
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    def test_io_generates_versioned_with_diary(self):
        self.client.force_authenticate(self.io)
        r1 = self.client.post(self.url)
        self.assertEqual(r1.status_code, status.HTTP_201_CREATED, r1.data)
        self.assertEqual(r1.data["doc_type"], "CHARGESHEET")
        self.assertEqual(r1.data["version"], 1)

        r2 = self.client.post(self.url)
        self.assertEqual(r2.data["version"], 2)
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="CHARGESHEET").count(), 2)
        gen_events = self.case.diary.filter(event_type="DOCUMENT_GENERATED")
        self.assertEqual(gen_events.count(), 2)
        self.assertIn("Purvani Chargesheet", gen_events.first().description)

    def test_legal_advisor_forbidden(self):
        self.client.force_authenticate(self.legal)
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="CHARGESHEET").count(), 0)


@override_settings(MEDIA_ROOT=_TEST_MEDIA)
class PanchanamaTests(APITestCase):
    """Week 2 Track B: Accused Panchanama (accused + witnesses + items)."""

    def setUp(self):
        self.io = User.objects.create_user(username="io_pan", password="pw", role="IO", police_station="PS-Central")
        self.legal = User.objects.create_user(username="legal_pan", password="pw", role="LEGAL")
        self.case = Case.objects.create(
            case_number="PAN-1", fir_number="FIR-PAN", fir_date="2026-06-01",
            police_station="PS-Central", incident_narrative="Break-in and theft.", io=self.io,
        )
        self.accused = Party.objects.create(case=self.case, role="ACCUSED", name="Ravi Kumar", parent_name="Mohan Kumar", address="12 MG Road")
        Party.objects.create(case=self.case, role="WITNESS", name="Sunil Verma", parent_name="Gopal", address="5 Nehru Road")
        Party.objects.create(case=self.case, role="WITNESS", name="Meena Joshi", parent_name="Hari", address="7 Tagore Lane")
        SeizedItem.objects.create(case=self.case, description="One mobile phone", quantity="1", seized_from=self.accused)
        SeizedItem.objects.create(case=self.case, description="Stolen gold chain", quantity="1", seized_from=self.accused)
        self.url = reverse("case-generate-panchanama", args=[self.case.id])

    def _all_text(self, case):
        from io import BytesIO

        from docx import Document as Docx

        from core.documents import build_panchanama_context_from_case, render_panchanama

        tpl = render_panchanama(build_panchanama_context_from_case(case))
        buf = BytesIO()
        tpl.save(buf)
        buf.seek(0)
        doc = Docx(buf)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)

    def test_render_lands_accused_witnesses_and_items(self):
        text = self._all_text(self.case)
        # Accused
        self.assertIn("Ravi Kumar", text)
        self.assertIn("12 MG Road", text)
        # Both witnesses (table loop)
        self.assertIn("Sunil Verma", text)
        self.assertIn("Meena Joshi", text)
        # Both items (table loop)
        self.assertIn("One mobile phone", text)
        self.assertIn("Stolen gold chain", text)
        # No leftover Jinja
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    def test_io_generates_versioned_with_diary(self):
        self.client.force_authenticate(self.io)
        r1 = self.client.post(self.url)
        self.assertEqual(r1.status_code, status.HTTP_201_CREATED, r1.data)
        self.assertEqual(r1.data["doc_type"], "PANCHANAMA")
        self.assertEqual(r1.data["version"], 1)
        r2 = self.client.post(self.url)
        self.assertEqual(r2.data["version"], 2)
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="PANCHANAMA").count(), 2)
        self.assertIn("Accused Panchanama", self.case.diary.filter(event_type="DOCUMENT_GENERATED").first().description)

    def test_legal_advisor_forbidden(self):
        self.client.force_authenticate(self.legal)
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="PANCHANAMA").count(), 0)


@override_settings(MEDIA_ROOT=_TEST_MEDIA)
class FaceIdFormTests(APITestCase):
    """Week 2 Track B: Face Identification Form (accused + photo placeholder)."""

    def setUp(self):
        self.io = User.objects.create_user(username="io_fid", password="pw", role="IO", police_station="PS-Central")
        self.legal = User.objects.create_user(username="legal_fid", password="pw", role="LEGAL")
        self.case = Case.objects.create(
            case_number="FID-1", fir_number="FIR-FID", fir_date="2026-06-01",
            police_station="PS-Central", incident_narrative="Theft.", io=self.io,
        )
        Party.objects.create(case=self.case, role="ACCUSED", name="Ravi Kumar", parent_name="Mohan Kumar", age=28, gender="Male", address="12 MG Road")
        self.url = reverse("case-generate-faceid", args=[self.case.id])

    def _all_text(self, case):
        from io import BytesIO

        from docx import Document as Docx

        from core.documents import build_faceid_context_from_case, render_faceid

        tpl = render_faceid(build_faceid_context_from_case(case))
        buf = BytesIO()
        tpl.save(buf)
        buf.seek(0)
        doc = Docx(buf)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)

    def test_render_lands_accused_and_photo_placeholder(self):
        text = self._all_text(self.case)
        self.assertIn("Ravi Kumar", text)
        self.assertIn("Mohan Kumar", text)
        self.assertIn("12 MG Road", text)
        # The photo placeholder box (no image field on Party yet).
        self.assertIn("AFFIX RECENT PHOTOGRAPH", text)
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)

    def test_io_generates_versioned_with_diary(self):
        self.client.force_authenticate(self.io)
        r1 = self.client.post(self.url)
        self.assertEqual(r1.status_code, status.HTTP_201_CREATED, r1.data)
        self.assertEqual(r1.data["doc_type"], "FACEID")
        self.assertEqual(r1.data["version"], 1)
        r2 = self.client.post(self.url)
        self.assertEqual(r2.data["version"], 2)
        self.assertIn("Face Identification Form", self.case.diary.filter(event_type="DOCUMENT_GENERATED").first().description)

    def test_legal_advisor_forbidden(self):
        self.client.force_authenticate(self.legal)
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(Document.objects.filter(case=self.case, doc_type="FACEID").count(), 0)


@override_settings(MEDIA_ROOT=_TEST_MEDIA)
class PdfExportTests(APITestCase):
    """Week 2 Track C: docx -> PDF export via LibreOffice headless."""

    def setUp(self):
        self.io = User.objects.create_user(username="io_pdf", password="pw", role="IO", police_station="PS-Central")
        self.case = Case.objects.create(
            case_number="PDF-1", fir_number="FIR-PDF", fir_date="2026-06-01",
            police_station="PS-Central", incident_narrative="Theft.", io=self.io,
        )
        Party.objects.create(case=self.case, role="ACCUSED", name="Ravi Kumar", parent_name="Mohan", address="12 MG Road")
        self.client.force_authenticate(self.io)
        self.url = reverse("case-generate-seizure", args=[self.case.id])

    def test_format_pdf_produces_valid_pdf(self):
        resp = self.client.post(self.url + "?format=pdf")
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED, resp.data)
        self.assertTrue(resp.data["pdf_url"].endswith(".pdf"))
        # Primary download follows the requested format.
        self.assertEqual(resp.data["download_url"], resp.data["pdf_url"])

        doc = Document.objects.get(case=self.case, doc_type="SEIZURE")
        self.assertTrue(doc.pdf_file)
        doc.pdf_file.open("rb")
        try:
            data = doc.pdf_file.read()
        finally:
            doc.pdf_file.close()
        # A valid, non-empty PDF.
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertGreater(len(data), 1000)

    def test_default_generation_has_no_pdf(self):
        resp = self.client.post(self.url)
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED, resp.data)
        self.assertIsNone(resp.data["pdf_url"])
        self.assertTrue(resp.data["download_url"].endswith(".docx"))
        doc = Document.objects.get(case=self.case, doc_type="SEIZURE")
        self.assertFalse(doc.pdf_file)
